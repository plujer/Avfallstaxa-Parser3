from openpyxl import Workbook
from openpyxl.worksheet.table import Table, TableStyleInfo
from excel_builder.schema import WorkbookSchemaScanner

def test_workbook_schema_scanner_detects_header_row_and_table(tmp_path):
    path = tmp_path / "master.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Taxepunkter"
    ws.append(["Rubrik"])
    ws.append([""])
    ws.append([""])
    ws.append([""])
    ws.append(["Paragraf", "Paragrafnamn", "Taxapunkt", "Variant", "Enhet", "Taxakod"])
    ws.append(["6.1.2", "", "Asbest", "", "kilogram", "EDP1"])
    table = Table(displayName="TaxepunkterTable", ref="A5:F6")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(table)
    wb.save(path)

    schema = WorkbookSchemaScanner().scan(path)
    sheet = schema.sheet("Taxepunkter")

    assert schema.sheet_count == 1
    assert sheet is not None
    assert sheet.detected_header_row == 5
    assert "TaxepunkterTable" in sheet.tables
    assert "Paragraf" in sheet.detected_headers

def test_workbook_schema_scanner_warns_missing_file(tmp_path):
    schema = WorkbookSchemaScanner().scan(tmp_path / "missing.xlsx")
    assert schema.sheet_count == 0
    assert schema.warnings
