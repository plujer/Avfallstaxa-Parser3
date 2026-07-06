from openpyxl import Workbook
from parser3.excel import MasterExcelReader, WorkbookProfiler


def test_workbook_profiler_finds_best_sheet(tmp_path):
    path = tmp_path / "master.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Info"
    ws["A1"] = "not relevant"
    ws2 = wb.create_sheet("Facit")
    ws2.append(["Paragraf", "Taxepunkt", "Variant", "Enhet", "Pris"])
    ws2.append(["6.1.2", "Asbest", "", "kilogram", "22,88"])
    wb.save(path)

    profile = WorkbookProfiler().profile(path)

    assert profile.best_sheet.sheet_name == "Facit"
    assert profile.best_sheet.detected_columns["section"] == 1
    assert profile.best_sheet.detected_columns["name"] == 2
    assert profile.best_sheet.detected_columns["price"] == 5


def test_master_excel_reader_reads_best_sheet(tmp_path):
    path = tmp_path / "master.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Start"
    ws.append(["x"])
    ws2 = wb.create_sheet("Taxor")
    ws2.append(["Paragraf", "Benämning", "Enhet", "Pris"])
    ws2.append(["2.1", "Fritidshus", "styck", "XX kr"])
    wb.save(path)

    rows = MasterExcelReader().read(path)

    assert len(rows) == 1
    assert rows[0].section == "2.1"
    assert rows[0].name == "Fritidshus"
    assert rows[0].unit == "styck"
    assert rows[0].price == "XX kr"
